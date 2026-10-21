"""Local package/content checks; does not claim native PowerPoint playback."""
from pathlib import Path
from zipfile import ZipFile
import sys,json,posixpath,hashlib
sys.path.insert(0,str(Path(__file__).parent/'.vendor'))
from lxml import etree as E
from PIL import Image
from pptx import Presentation

OUT=Path(__file__).parent;ROOT=OUT.parent
P='http://schemas.openxmlformats.org/presentationml/2006/main'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS={'p':P,'a':A,'w':W}

def package(path):
    with ZipFile(path) as z:
        assert z.testzip() is None
        names=set(z.namelist())
        for name in names:
            if name.endswith(('.xml','.rels')):
                root=E.fromstring(z.read(name))
                if name.endswith('.rels'):
                    base=posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        if rel.get('TargetMode')=='External':continue
                        target=posixpath.normpath(posixpath.join(base,rel.get('Target'))).lstrip('/')
                        assert target in names,(path,name,target)
    return True

def main():
    deck=OUT/'CyberForge - Thesis Defence.pptx';doc=ROOT/'Che Haggai Ngha Kang.docx'
    for p in [deck,doc,OUT/'Che Haggai Ngha Kang - Original backup.docx']:package(p)
    prs=Presentation(deck);assert len(prs.slides)==15
    expected=['CyberForge','Brief introduction','Problem statement','Research questions','Research objectives','Brief literature review']+['Methods and materials used']*4+['Results']*4+['Conclusion']
    for i,(slide,title) in enumerate(zip(prs.slides,expected),1):
        texts=[sh.text for sh in slide.shapes if sh.has_text_frame]
        assert title in texts,('required guide heading missing',i,title)
        assert not any(text.startswith('CYBERFORGE  /') and 'CURRENT SYSTEM' in text for text in texts),('off-guide section',i)
    builds=0
    for s in prs.slides:
        ids={str(sh.shape_id) for sh in s.shapes}
        assert s._element.find('{'+P+'}transition') is not None
        timing=s._element.find('{'+P+'}timing');assert timing is not None
        tids=[n.get('id') for n in timing.findall('.//p:cTn',NS)];assert len(tids)==len(set(tids))
        for target in timing.findall('.//p:spTgt',NS):assert target.get('spid') in ids
        builds+=len(timing.findall('.//p:animEffect',NS))
        assert 'Source:' in s.notes_slide.notes_text_frame.text
    for name in ['orbit.gif','signal.gif']:
        im=Image.open(OUT/'assets'/name);assert im.n_frames>=40 and im.info.get('loop')==0
    for box in json.loads((OUT/'layout-audit.json').read_text(encoding='utf-8')):
        assert box['x']>=0 and box['y']>=0 and box['x']+box['w']<=16.01 and box['y']+box['h']<=9.01,box
    with ZipFile(doc) as z:
        r=E.fromstring(z.read('word/document.xml'));text=' '.join(r.itertext())
        for phrase in ['Tauri 2','Node.js 24','HistGradientBoostingClassifier','IsolationForest','Qwen','MCP','Appendix C','historical','4.05']:
            assert phrase in text,phrase
        assert r.findall('.//w:oMath',NS)==[] # Office math uses its own namespace, checked separately below.
        assert r.find('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath') is not None
        body=r.find('w:body',NS)
        report=json.loads((OUT/'thesis-revision-log.json').read_text(encoding='utf-8'))
        assert len(report['paragraphRevisions'])>=110
        assert hashlib.sha256((OUT/'Che Haggai Ngha Kang - Original backup.docx').read_bytes()).hexdigest()==report['originalSHA256']
    # Every original historic result table is byte-for-byte XML-equivalent except
    # Table 4.8's deliberately corrected latency target/outcome row.
    with ZipFile(OUT/'Che Haggai Ngha Kang - Original backup.docx') as z:
        original=E.fromstring(z.read('word/document.xml')).find('w:body',NS)
        for i in [514,517,529,532,544,555,565]:
            assert E.tostring(original[i])==E.tostring(body[i]),('historical table changed',i)
    result={'slides':15,'nativeEntranceEffects':builds,'animatedGIFs':2,'transitions':15,'speakerNotes':15,'revisedParagraphs':len(report['paragraphRevisions']),'historicalTablesPreserved':7,'packageRelationships':'valid','slideBounds':'valid','nativeOfficePlayback':'not available in this environment; package structure and generated previews checked'}
    (OUT/'validation-report.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
