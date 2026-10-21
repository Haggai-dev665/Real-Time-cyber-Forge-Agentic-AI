"""Build editable 15-slide thesis deck, animation assets and review previews.

Uses the thesis content and audited local source; no generated experimental data.
PowerPoint shapes are mirrored into PNG/PDF previews for visual inspection.
"""
import sys, math, json, textwrap
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent / '.vendor'))
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree as E

OUT=Path(__file__).parent
ASSETS=OUT/'assets'; ASSETS.mkdir(exist_ok=True)
PREV=OUT/'previews'; PREV.mkdir(exist_ok=True)
BG='1F1A1B'; PANEL='2A2425'; CREAM='EFE9DD'; ORANGE='F69D39'; SAGE='869B7E'; MUTED='C2B8AA'; INK='0B0908'; LINE='51474A'
HEAD='Bahnschrift'; BODY='Calibri'; MONO='Consolas'
FONTDIR=Path('C:/Windows/Fonts')
FONTFILES={HEAD:'bahnschrift.ttf',BODY:'calibri.ttf',MONO:'consola.ttf'}
SCALE=100
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
prs.core_properties.title='CyberForge — Thesis Defence'
prs.core_properties.author='Che Haggai Ngha Kang'
prs.core_properties.subject='Historical research results and current implementation review, 18 October 2026'
prs.core_properties.keywords='CyberForge, thesis, agentic AI, web threat detection, University of Bamenda'
slides=[]; all_notes=[]; boxes=[]

def rgb(c):return RGBColor.from_string(c)
def q(tag):return OxmlElement(tag)
def font(sz,family=BODY):return ImageFont.truetype(str(FONTDIR/FONTFILES[family]),round(sz*SCALE/72))
def wrapped(text, width, sz, family=BODY):
    f=font(sz,family);lines=[]
    for raw in text.split('\n'):
        line=''
        for word in raw.split(' '):
            candidate=(line+' '+word).strip()
            if line and f.getlength(candidate)>width*SCALE:lines.append(line);line=word
            else:line=candidate
        lines.append(line)
    return lines

def motion_assets():
    frames=[]
    for k in range(48):
        im=Image.new('RGB',(600,600),'#'+BG);d=ImageDraw.Draw(im)
        cx=cy=300
        for rad in [120,180,245]:d.ellipse((cx-rad,cy-rad,cx+rad,cy+rad),outline='#'+LINE,width=2)
        for j in range(8):
            a=j*math.tau/8-math.pi/2;x=cx+245*math.cos(a);y=cy+245*math.sin(a)
            d.line((cx+115*math.cos(a),cy+115*math.sin(a),x,y),fill='#'+LINE,width=2)
            d.ellipse((x-7,y-7,x+7,y+7),fill='#'+(ORANGE if j%2 else SAGE))
        a=k/48*math.tau
        for rad,delta in [(180,0),(245,2.4)]:
            x=cx+rad*math.cos(a+delta);y=cy+rad*math.sin(a+delta)
            d.arc((cx-rad,cy-rad,cx+rad,cy+rad),int((a+delta)*180/math.pi)-45,int((a+delta)*180/math.pi),fill='#'+ORANGE,width=4)
            d.ellipse((x-6,y-6,x+6,y+6),fill='#'+CREAM)
        shield=[(300,185),(390,220),(375,330),(300,405),(225,330),(210,220)]
        d.polygon(shield,fill='#'+PANEL,outline='#'+ORANGE,width=5)
        d.line([(260,292),(291,326),(345,259)],fill='#'+CREAM,width=10)
        frames.append(im)
    frames[0].save(ASSETS/'orbit.gif',save_all=True,append_images=frames[1:],duration=75,loop=0,optimize=True)
    frames=[]
    for k in range(40):
        im=Image.new('RGB',(1440,20),'#'+BG);d=ImageDraw.Draw(im)
        d.line((0,10,1440,10),fill='#'+LINE,width=1)
        for j in range(5):
            x=(k*36+j*288)%1440
            for t in range(45):
                a=(1-t/45)*.85;bg=tuple(int(BG[i:i+2],16) for i in (0,2,4));fg=tuple(int(ORANGE[i:i+2],16) for i in (0,2,4));c=tuple(round(bg[i]*(1-a)+fg[i]*a) for i in range(3))
                d.line(((x-t)%1440,9,(x-t)%1440,11),fill=c,width=1)
        frames.append(im)
    frames[0].save(ASSETS/'signal.gif',save_all=True,append_images=frames[1:],duration=90,loop=0,optimize=True)

class Slide:
    def __init__(self,section,title,subtitle,source,notes):
        self.s=prs.slides.add_slide(prs.slide_layouts[6]); self.n=len(prs.slides)
        self.s.background.fill.solid();self.s.background.fill.fore_color.rgb=rgb(BG)
        self.im=Image.new('RGB',(1600,900),'#'+BG);self.draw=ImageDraw.Draw(self.im)
        self.groups={};self.texts=[]
        self.text(.7,.32,12,.3,'CYBERFORGE  /  '+section.upper(),11,ORANGE,MONO)
        self.text(14,.32,1.3,.3,f'{self.n:02d} / 15',11,MUTED,MONO,align='right')
        self.picture(ASSETS/'signal.gif',.7,.78,14.6,.12)
        if title:self.text(.7,1.12,14.1,.8,title,36,CREAM,HEAD,group=0)
        if subtitle:self.text(.73,2.02,14,.55,subtitle,17,MUTED,group=0)
        self.rect(.7,8.35,14.6,.015,LINE)
        self.text(.7,8.48,13.6,.3,source,10,MUTED)
        self.rect(.7,8.92,14.6*self.n/15,.04,ORANGE)
        self.s.notes_slide.notes_text_frame.text=notes+'\n\nSource: '+source+'\n\nPresentation control: slide builds run automatically; advance manually after the explanation. Animated GIF accents loop in Slide Show. Metrics shown on results slides belong to the original thesis experiment.'
        all_notes.append({'slide':self.n,'title':title or 'CyberForge','source':source,'notes':notes})
        slides.append(self)
    def track(self,shape,group):
        if group is not None:self.groups.setdefault(group,[]).append(shape.shape_id)
        return shape
    def rect(self,x,y,w,h,color=PANEL,stroke=None,group=None,round=False):
        sh=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
        if stroke:sh.line.color.rgb=rgb(stroke);sh.line.width=Pt(.8)
        else:sh.line.fill.background()
        if round:
            sh.adjustments[0]=.08
            self.draw.rounded_rectangle((x*100,y*100,(x+w)*100,(y+h)*100),radius=12,fill='#'+color,outline='#'+stroke if stroke else None,width=1)
        else:self.draw.rectangle((x*100,y*100,(x+w)*100,(y+h)*100),fill='#'+color,outline='#'+stroke if stroke else None)
        return self.track(sh,group)
    def text(self,x,y,w,h,text,sz=22,color=CREAM,family=BODY,group=None,align='left'):
        lines=wrapped(text,w,sz,family); lineheight=sz/72*1.16
        if len(lines)*lineheight>h+.06: raise ValueError(f'Slide {self.n} text overflow: {text[:60]} ({len(lines)*lineheight:.2f} > {h})')
        sh=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=family;p.font.size=Pt(sz);p.font.color.rgb=rgb(color)
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.16
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            px=x*100;f=font(sz,family)
            if align=='center':px+=(w*100-f.getlength(line))/2
            if align=='right':px+=w*100-f.getlength(line)
            self.draw.text((px,(y+i*lineheight)*100),line,font=f,fill='#'+color,anchor='lt')
        self.texts.append(text);boxes.append({'slide':self.n,'x':x,'y':y,'w':w,'h':h,'text':text})
        return self.track(sh,group)
    def line(self,x1,y1,x2,y2,color=SAGE,group=None,arrow=False,width=2):
        sh=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
        self.draw.line((x1*100,y1*100,x2*100,y2*100),fill='#'+color,width=max(1,round(width*100/72)))
        if arrow:
            end=q('a:tailEnd');end.set('type','triangle');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(end)
            a=math.atan2(y2-y1,x2-x1);p=[(x2*100,y2*100),((x2-.12*math.cos(a-.5))*100,(y2-.12*math.sin(a-.5))*100),((x2-.12*math.cos(a+.5))*100,(y2-.12*math.sin(a+.5))*100)];self.draw.polygon(p,fill='#'+color)
        return self.track(sh,group)
    def picture(self,path,x,y,w,h,group=None):
        sh=self.s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h));im=Image.open(path).convert('RGBA').resize((round(w*100),round(h*100)),Image.Resampling.LANCZOS);self.im.paste(im,(round(x*100),round(y*100)),im)
        return self.track(sh,group)
    def card(self,x,y,w,h,kicker,title,body,group=1,accent=ORANGE):
        self.rect(x,y,w,h,PANEL,LINE,group,True);self.rect(x,y,.045,h,accent,group=group)
        self.text(x+.23,y+.2,w-.46,.3,kicker,11,accent,MONO,group)
        self.text(x+.23,y+.72,w-.46,.9,title,23,CREAM,HEAD,group)
        self.text(x+.23,y+1.78,w-.46,h-1.9,body,18,MUTED,BODY,group)
    def node(self,x,y,w,h,title,sub='',group=1,color=ORANGE):
        self.rect(x,y,w,h,PANEL,LINE,group,True);self.rect(x,y,.045,h,color,group=group)
        self.text(x+.2,y+.16,w-.4,.42,title,19,CREAM,HEAD,group)
        if sub:self.text(x+.2,y+.66,w-.4,h-.75,sub,14,MUTED,BODY,group)
    def animate(self):
        root=self.s._element
        tr=q('p:transition');tr.set('spd','med');tr.set('advClick','1')
        effect=q('p:fade');tr.append(effect);root.append(tr)
        timing=q('p:timing');tn=q('p:tnLst');timing.append(tn);par=q('p:par');tn.append(par)
        ct=q('p:cTn');ct.set('id','1');ct.set('dur','indefinite');ct.set('restart','never');ct.set('nodeType','tmRoot');par.append(ct)
        children=q('p:childTnLst');ct.append(children)
        seq=q('p:seq');seq.set('concurrent','1');seq.set('nextAc','seek');children.append(seq)
        main=q('p:cTn');main.set('id','2');main.set('dur','indefinite');main.set('nodeType','mainSeq');seq.append(main)
        mainkids=q('p:childTnLst');main.append(mainkids)
        kickoff=q('p:par');mainkids.append(kickoff)
        kc=q('p:cTn');kc.set('id','3');kc.set('fill','hold');kickoff.append(kc)
        start=q('p:stCondLst');cond=q('p:cond');cond.set('delay','0');start.append(cond);kc.append(start)
        builds=q('p:childTnLst');kc.append(builds);counter=4
        # A standard on-slide-load sequence: staggered automatic builds, no extra clicks.
        for group,ids in sorted(self.groups.items()):
            for sid in ids:
                shape=next(sh for sh in self.s.shapes if sh.shape_id==sid)
                wipe=shape.shape_type==1 and shape.height<Inches(.4) and shape.width>Inches(.8)
                outer=q('p:par');builds.append(outer);c=q('p:cTn');c.set('id',str(counter));counter+=1;c.set('fill','hold');outer.append(c)
                conds=q('p:stCondLst');cond=q('p:cond');cond.set('delay',str(180+group*480));conds.append(cond);c.append(conds)
                ks=q('p:childTnLst');c.append(ks);p=q('p:par');ks.append(p)
                ec=q('p:cTn');ec.set('id',str(counter));counter+=1;ec.set('presetID','22' if wipe else '10');ec.set('presetClass','entr');ec.set('presetSubtype','2' if wipe else '0');ec.set('fill','hold');ec.set('nodeType','withEffect');p.append(ec)
                st=q('p:stCondLst');co=q('p:cond');co.set('delay','0');st.append(co);ec.append(st)
                effects=q('p:childTnLst');ec.append(effects)
                setel=q('p:set');effects.append(setel);bh=q('p:cBhvr');setel.append(bh)
                tt=q('p:cTn');tt.set('id',str(counter));counter+=1;tt.set('dur','1');tt.set('fill','hold');bh.append(tt)
                target=q('p:tgtEl');sp=q('p:spTgt');sp.set('spid',str(sid));target.append(sp);bh.append(target)
                attrs=q('p:attrNameLst');a=q('p:attrName');a.text='style.visibility';attrs.append(a);bh.append(attrs)
                to=q('p:to');sv=q('p:strVal');sv.set('val','visible');to.append(sv);setel.append(to)
                anim=q('p:animEffect');anim.set('transition','in');anim.set('filter','wipe(right)' if wipe else 'fade');effects.append(anim)
                bh=q('p:cBhvr');anim.append(bh);tt=q('p:cTn');tt.set('id',str(counter));counter+=1;tt.set('dur','500');bh.append(tt)
                target=q('p:tgtEl');sp=q('p:spTgt');sp.set('spid',str(sid));target.append(sp);bh.append(target)
        for name,event in [('prevCondLst','onPrev'),('nextCondLst','onNext')]:
            cs=q('p:'+name);co=q('p:cond');co.set('evt',event);co.set('delay','0');tg=q('p:tgtEl');tg.append(q('p:sldTgt'));co.append(tg);cs.append(co);seq.append(cs)
        root.append(timing)
        self.im.save(PREV/f'{self.n:02d}.png')

def build():
    motion_assets()
    s=Slide('Master’s thesis defence','','','The University of Bamenda  /  COLTECH  /  Computer Engineering',
        'Introduce the study title and candidate. Supervisor: Prof. Onabid Mathias Akong. Registration UBA24PP134. The original dissertation is dated June 2026; this presentation incorporates the source review of 18 October 2026. Distinguish the historical experiment from current implementation evidence throughout.')
    s.text(.75,1.5,10,1.25,'CyberForge',70,CREAM,HEAD,0)
    s.text(.8,3.0,9.7,1.7,'A Real-Time Agentic AI System\nfor Web Threat Detection\nand User Protection',30,CREAM,HEAD,1)
    s.rect(.8,5.04,1.1,.055,ORANGE,group=1)
    s.text(.8,5.5,9,1,'CHE HAGGAI NGHA KANG',24,ORANGE,HEAD,2)
    s.text(.8,6.15,9,1.1,'UBA24PP134  •  Master’s in Computer Engineering\nInformation Technology and Cyber Security',18,MUTED,group=2)
    s.text(.8,7.42,10,.5,'Supervisor: Prof. ONABID MATHIAS AKONG',17,CREAM,group=3)
    s.picture(ASSETS/'orbit.gif',10.55,2,5,5,1)
    s.text(10.8,7.2,4.3,.7,'RESEARCH + CURRENT SYSTEM\nSource review · 18 October 2026',12,SAGE,MONO,3,align='center')

    s=Slide('01 / Brief introduction','Brief introduction','CyberForge: real-time web-threat detection through cooperating specialist agents.','Thesis §§1.1, 1.6; revised §3.2',
        'Explain phishing, malicious domains, redirects and credential theft as web-context risks. The system aims to connect evidence with a usable explanation. Current monitoring is consent-based and platform-dependent; do not claim every page is captured or blocked before it loads. This is not a full endpoint or network intrusion-detection suite.')
    for i,(t,b) in enumerate([('Observe','Submitted URL or permitted browsing signal'),('Analyse','Models + page signals + historical context'),('Explain','Verdict, supporting evidence and limitations'),('Respond','Warning or supported, approved action')]):
        x=.75+i*3.85;s.node(x,3.2,3.35,2,t,b,i+1)
        if i<3:s.line(x+3.38,4.2,x+3.8,4.2,ORANGE,i+1,True)
    s.text(.8,6.1,13.6,1.3,'Research idea: combine specialised analysers\ninto one inspectable security workflow.',29,CREAM,HEAD,5)

    s=Slide('02 / Problem statement','Problem statement','Existing approaches struggle with false alarms, unclear verdicts and fragmented context.','Thesis §1.2; original calibration described in §3.7.2',
        'The original scoring baseline added +0.5 and overflagged safe calibration examples. Present this as a development defect on a small set, not a population-wide model result. The other problems are unexplained classifier output and fragmented context. Enterprise SOAR motivates integration but this study does not benchmark commercial suites.')
    s.card(.75,2.95,4.6,4.45,'01 / TRUST','Too many false alarms','Biased scores interrupt safe browsing and encourage alert fatigue.',1)
    s.card(5.7,2.95,4.6,4.45,'02 / EVIDENCE','A score without a story','A single verdict does not explain indicators, behaviour or attack techniques.',2,SAGE)
    s.card(10.65,2.95,4.6,4.45,'03 / CONTINUITY','Disconnected context','Past observations and threat intelligence are difficult to reuse consistently.',3)

    s=Slide('03 / Research questions','Research questions','General question: real-time, explainable web-threat analysis across interfaces.','Thesis §1.3; validity qualifications in revised §§3.9 and 4.9',
        'State the four specific questions in sequence. The first asks about the smallest sufficient agent set; the thesis does not supply an ablation that resolves minimality. Memory reuse is implemented context, not evidence that model weights learn online. The evaluation supports an initial integrated-system result rather than complete answers to every causal question.')
    questions=[('01','Decomposition','How should URL analysis be divided, and what agent set is sufficient?'),('02','Integration','How can URL, content, IOC, ATT&CK and reputation signals agree?'),('03','Decision','How can aggregation reduce false alarms without hiding strong threats?'),('04','Context','How can past browsing observations support future assessment?')]
    for i,(num,t,b) in enumerate(questions):
        y=2.92+i*1.22;s.text(.8,y,.8,.65,num,31,ORANGE,MONO,i+1);s.text(1.9,y,3,.5,t,25,CREAM,HEAD,i+1);s.text(5.1,y,9.8,.85,b,21,MUTED,group=i+1);s.line(1.9,y+1.0,15.1,y+1.0,LINE,i+1,width=.6)

    s=Slide('04 / Research objectives','Research objectives','Build and evaluate an explainable, multi-platform web-security system.','Thesis §1.4; source-reviewed boundaries in Appendix C',
        'Present the general objective, then four specific objectives: coordinated tasks; multiple analytical perspectives; false-positive-aware aggregation; contextual external intelligence. Historical context is a supporting capability. Avoid claiming complete cross-platform protection or fully autonomous planning: current interfaces have different supported subsets.')
    for i,(title,body) in enumerate([('Coordinate agents','Decompose URL analysis into ordered, parallel tasks.'),('Unify evidence','Combine model, structural, content and behavioural signals.'),('Control false alerts','Balance aggregated evidence with strong threat indicators.'),('Add threat context','Link findings with attack techniques and external intelligence.')]):
        x=.75+(i%2)*7.5;y=2.95+(i//2)*2.45;s.node(x,y,7.05,2.05,title,body,i+1,color=SAGE if i%2 else ORANGE)

    s=Slide('05 / Brief literature review','Brief literature review','Prior work motivates hybrid detection, structured explanations and agent coordination.','Thesis Chapter 2; full bibliographic details are in the speaker notes',
        'Literature: Khonji, Iraqi & Jones (2013), Phishing detection: A literature survey, IEEE Communications Surveys & Tutorials 15(4), 2091–2121. Le, Pham, Sahoo & Hoi (2018), URLNet, arXiv:1802.03162. Devlin et al. (2019), BERT, NAACL-HLT, 4171–4186. Strom et al. (2018), MITRE ATT&CK: Design and philosophy. Wooldridge (2009), An introduction to multiagent systems, 2nd ed. Yao et al. (2023), ReAct, ICLR. Sommer & Paxson (2010), Outside the Closed World, IEEE S&P, 305–316. These sources motivate design; they do not prove CyberForge outperforms those systems.')
    for i,(label,title,body) in enumerate([('Khonji et al. · 2013','Hybrid signals','URL, page content and reputation offer complementary evidence.'),('Le · 2018 / Devlin · 2019','Learned representations','URLNet and BERT reduce reliance on manually selected patterns.'),('Strom et al. · 2018','Structured explanation','ATT&CK gives findings a shared vocabulary of attack techniques.'),('Wooldridge · 2009\nYao et al. · 2023','Agent coordination','Specialisation and tool-based reasoning motivate modular workflows.')]):
        x=.75+i*3.85;s.rect(x,3,3.35,3.75,PANEL,LINE,i+1,True);s.text(x+.2,3.2,2.95,.65,label,12,SAGE,MONO,i+1);s.text(x+.2,4.1,2.95,1,title,22,CREAM,HEAD,i+1);s.text(x+.2,5.15,2.95,1.4,body,17,MUTED,group=i+1)
    s.text(.8,7.2,14.4,.8,'Research gap: reproducible integration, meaningful failures and usable evidence.',23,ORANGE,HEAD,5)

    s=Slide('06 / Methods and materials used','Methods and materials used','1 of 4 · Software materials and system architecture','Revised §3.2 and Appendix C; Cargo.toml, backend/package.json, hf_space_deploy/',
        'The desktop is Tauri 2/Rust, not Electron. Backend runtime is Node.js 24/Express 4; the canonical Space container uses Python 3.11/FastAPI with no Gradio UI. Shared core serves a limited CLI and grant-scoped stdio MCP surface. MongoDB owns control-plane records, while Appwrite and local/optional stores have separate roles. Mobile source is React Native/Expo; native-capture release acceptance is not asserted. These are source specifications, not a deployment availability certificate.')
    rows=[('CLIENTS','Tauri 2 desktop  •  Rust CLI  •  MCP stdio  •  Expo companion'),('CONTROL','Rust core + keyring + consent  →  Node.js 24 / Express APIs'),('ANALYSIS','Specialist DAG / durable runs  →  Python 3.11 / FastAPI on Hugging Face'),('DATA','MongoDB  •  Appwrite integrations  •  local context  •  configured OTX / caches')]
    for i,(label,text) in enumerate(rows):
        y=2.85+i*1.14;s.rect(.8,y,2.0,.84,ORANGE if i%2==0 else SAGE,group=i+1,round=True);s.text(.95,y+.24,1.7,.4,label,15,INK,MONO,i+1,align='center');s.node(3.1,y,12,.84,text,'',i+1)
        if i<3:s.line(8.8,y+.86,8.8,y+1.1,ORANGE,i+1,True)
    s.text(.85,7.7,14.4,.45,'Identity, consent and evidence cross every layer; model availability is reported separately.',18,MUTED,group=5)

    s=Slide('06 / Methods and materials used','Methods and materials used','2 of 4 · Eight-agent analysis pipeline and coordination method','Thesis Table 3.1; backend/src/services/agentOrchestrator.js',
        'Stage 1 launches four tasks: classifier, DGA, scraper and memory. Stage 2 consumes collected page evidence for IOC extraction, behavioural rules and MITRE mapping. Stage 3 enriches IOCs through OTX. The Reporter follows the eight specialists, so it is not counted among the eight. Multiple agents produce evidence, but the original study does not contain an ablation proving the smallest effective set. The CLI/MCP versioned detector operation does not automatically run this entire page DAG.')
    for x,label in [(1.0,'STAGE 1 · COLLECT'),(5.3,'STAGE 2 · INSPECT'),(9.5,'STAGE 3 · ENRICH')]:s.text(x,2.75,4,.4,label,13,SAGE,MONO,1)
    for i,t in enumerate(['URL classifier · BERT','DGA detector · rules','Web scraper','Memory reader']):s.node(.85,3.35+i*.87,3.5,.68,t,group=1)
    for i,t in enumerate(['IOC extractor','Behavioural analyser','MITRE ATT&CK mapper']):s.node(5.05,3.8+i*1.04,3.6,.76,t,group=2)
    s.node(9.4,4.25,2.5,1.6,'Threat intel','OTX reputation',3)
    s.node(12.65,4.25,2.5,1.6,'Reporter','Verdict + evidence',4)
    s.line(4.4,4.9,4.95,4.9,ORANGE,2,True);s.line(8.7,4.9,9.3,4.9,ORANGE,3,True);s.line(11.95,4.9,12.55,4.9,ORANGE,4,True)
    s.line(2.6,7.45,13.9,7.45,SAGE,5);s.line(13.9,5.9,13.9,7.45,SAGE,5);s.line(2.6,7.45,2.6,6.75,SAGE,5,True)
    s.text(4.1,7.65,8.2,.4,'Persisted context supports later visits — it does not retrain the models.',16,MUTED,group=5,align='center')

    s=Slide('06 / Methods and materials used','Methods and materials used','3 of 4 · Learning techniques and model-training phases','Revised Appendix C.3–C.4; hf_space_deploy/trainer.py, app.py, url_features.py',
        'Three HistGradientBoostingClassifier models learn supervised URL labels for phishing, malware-serving URLs and web-attack payloads. IsolationForest is fitted on benign examples for unsupervised anomaly scoring. BERT and Qwen are externally pretrained inference, not locally fine-tuned here. DGA and IOC logic are rules. Trainer phases: collect/clean; shared features and domain-grouped split; fit; threshold selection and held-out acceptance; publish accepted artefacts. Defaults include 20% domain-held-out test and a further 10% of training for threshold choice. Training shares the serving process, so isolation, durable jobs and release rollback remain work. No reinforcement learning is established.')
    for i,(t,b) in enumerate([('SUPERVISED','3 gradient-boosted\nURL classifiers'),('UNSUPERVISED','IsolationForest\ntrained on benign URLs'),('PRETRAINED + RULES','BERT / Qwen inference\nDGA / IOC heuristics')]):
        x=.8+i*5.0;s.node(x,2.95,4.6,1.85,t,b,i+1)
    stages=['Collect\n& clean','Features\n& split','Fit\nmodels','Evaluate\n& gate','Publish\naccepted']
    for i,t in enumerate(stages):
        x=.85+i*2.98;s.rect(x,5.55,2.5,1.05,SAGE if i==3 else PANEL,LINE,4+i,True);s.text(x+.12,5.72,2.26,.8,t,18,INK if i==3 else CREAM,HEAD,4+i,align='center')
        if i<4:s.line(x+2.55,6.07,x+2.9,6.07,ORANGE,4+i,True)
    s.text(.85,7.25,14.3,.7,'Domain-held-out evaluation  •  Shared training/serving features  •  No automatic online weight updates',19,ORANGE,group=9)

    s=Slide('06 / Methods and materials used','Methods and materials used','4 of 4 · Evaluation data, calibration and testing procedure','Thesis §§3.7–3.9, Table 4.3; revised Appendix C.5',
        'Original development used seven calibration URLs and 30 DGA-style/legitimate labels. The reported evaluation used 300 URLs: 150 benign, 100 phishing and 50 malware/DGA. The thesis describes 100 consecutive requests as its latency protocol but supplies no raw traces here. Do not imply that current training corpus sizes were the historical evaluation size. New source tests validate permissions and contracts, not current detection accuracy. No comparative ablation, confidence intervals or reproducible per-URL manifest are supplied in the original document.')
    s.text(.85,2.9,5,1.1,'300 URLs',54,ORANGE,HEAD,1)
    s.text(.9,4.05,4.4,1.25,'Historical classification\nevaluation set',24,CREAM,HEAD,1)
    for i,(v,t,c) in enumerate([(150,'benign',SAGE),(100,'phishing',ORANGE),(50,'malware / DGA',MUTED)]):
        y=3.03+i*1.06;s.text(6.2,y,1.4,.55,str(v),27,c,MONO,2+i);s.text(8,y,6,.55,t,24,CREAM,group=2+i);s.rect(6.2,y+.67,8.8*v/150,.13,c,group=2+i)
    s.rect(.85,6.25,14.3,1.5,PANEL,LINE,5,True)
    s.text(1.1,6.5,6.7,.8,'Development calibration\n7 URL examples + 30 domain labels',18,MUTED,group=5)
    s.text(8.1,6.5,6.7,.8,'Current source verification\nContracts, access, failure and UI behaviour',18,MUTED,group=5)

    s=Slide('07 / Results','Results','1 of 4 · Detection performance: 284 of 300 URLs classified correctly','Thesis Tables 4.3–4.4; confusion matrix derived directly from Table 4.3',
        'Actual harmful URLs: 137 detected and 13 missed. Actual benign URLs: 147 correctly left alone and three false alarms. Accuracy=(137+147)/300=94.7%; precision=137/140=97.9%; recall=137/150=91.3%; specificity=147/150=98.0%; F1=274/290=94.5%; false-positive rate=3/150=2.0%. The counts support these metrics but not a causal claim that the multi-agent design is superior to a single-model baseline. Category-level detection is 92/100 phishing and 45/50 malware/DGA.')
    s.text(1.0,2.85,2.9,.4,'PREDICTED THREAT',11,MUTED,MONO,1,align='center')
    s.text(4.15,2.85,2.9,.4,'PREDICTED BENIGN',11,MUTED,MONO,1,align='center')
    for x,y,n,label,c in [(1.0,3.45,'137','True positives',SAGE),(4.15,3.45,'13','False negatives',ORANGE),(1.0,5.65,'3','False positives',ORANGE),(4.15,5.65,'147','True negatives',SAGE)]:
        s.rect(x,y,2.9,1.85,PANEL,LINE,1 if y<5 else 2,True);s.text(x+.2,y+.2,2.5,.8,n,42,c,HEAD,1 if y<5 else 2);s.text(x+.2,y+1.22,2.5,.4,label,15,MUTED,group=1 if y<5 else 2)
    s.text(1.05,3.2,5.8,.22,'ACTUAL HARMFUL',10,ORANGE,MONO,1);s.text(1.05,5.37,5.8,.22,'ACTUAL BENIGN',10,SAGE,MONO,2)
    for i,(v,l) in enumerate([('94.7%','Accuracy'),('97.9%','Precision'),('91.3%','Recall'),('2.0%','False-positive rate')]):
        x=8.2+(i%2)*3.7;y=3.3+(i//2)*2.22;s.text(x,y,3.4,.92,v,43,ORANGE if i==3 else CREAM,HEAD,3+i);s.text(x,y+1.0,3.4,.6,l,18,MUTED,group=3+i)

    s=Slide('07 / Results','Results','2 of 4 · Processing speed: 4.05-second mean end-to-end analysis','Thesis Tables 4.1–4.2 and §4.6; original warm target in §3.9',
        'Means: URL/context 850 ms; content 1,200; intelligence 600; reporter 1,400; total 4,050. The reporter is the largest listed stage. The strict original warm target is 4,000 ms, so the mean misses it by 50 ms; the later summary used a looser approximately five-second target. The slides retain the stricter interpretation. Scraping is grouped under content in the timing table but launches in Stage 1 in the architectural DAG. Cached repeats below 100 ms and about 15% repeated queries are historical claims, not new measurements. No distribution/p95 can be derived from the supplied mean.')
    labels=[('URL / context',850,SAGE),('Content analysis',1200,CREAM),('Threat intelligence',600,SAGE),('Reporter',1400,ORANGE)]
    for i,(lab,v,c) in enumerate(labels):
        y=3.0+i*1.02;s.text(.85,y,3.5,.5,lab,20,CREAM,group=i+1);s.rect(4.5,y+.1,6.5*v/1400,.3,c,group=i+1);s.text(11.3,y,1.8,.5,f'{v:,}',22,c,MONO,i+1)
    s.text(13.2,3.0,2.1,1,'4.05 s',33,ORANGE,HEAD,5)
    s.text(13.2,4.1,2.1,1.7,'Mean total\n\n<100 ms\ncached repeat',18,MUTED,group=5)
    s.rect(.85,7.25,14.3,.67,PANEL,LINE,6,True);s.text(1.1,7.4,13.8,.4,'Strict warm target: 4.00 s  →  measured mean: 0.05 s above target.',19,ORANGE,group=6)

    s=Slide('07 / Results','Results','3 of 4 · Explanation quality and the user-facing alert','Thesis Table 4.6 and Fig. 4.3 (original alert screenshot)',
        'Historical manual ratings: factual accuracy 9.5/10, readability 9.0 and usefulness of advice 8.8. The original document does not state rater count, blinding or agreement, so call these descriptive ratings. They evaluated the former Gemini/Mistral setup, not current Qwen quality. The alert screenshot is reproduced from the supplied thesis and is an interface illustration, not additional performance evidence. Current actions must distinguish proposals, approvals and applied receipts.')
    for i,(lab,v) in enumerate([('Factual accuracy',9.5),('Readability',9.0),('Useful advice',8.8)]):
        y=3.15+i*1.18;s.text(.9,y,6,.45,lab,22,CREAM,group=i+1);s.rect(.9,y+.62,6.3,.16,LINE,group=i+1);s.rect(.9,y+.62,6.3*v/10,.16,SAGE if i==1 else ORANGE,group=i+1);s.text(7.55,y,2,.6,f'{v:.1f}/10',25,ORANGE,MONO,i+1)
    s.text(.9,7.0,8.5,.85,'Human-rated, descriptive results.\nNot a blinded user study or a current-model score.',17,MUTED,group=4)
    image=Image.open(ASSETS/'image11.png');w=3.6;h=w*image.height/image.width
    s.picture(ASSETS/'image11.png',10.65,2.85,w,h,3)
    s.text(10.25,7.55,4.5,.5,'Original thesis interface illustration',12,MUTED,group=4,align='center')

    s=Slide('07 / Results','Results','4 of 4 · Findings in relation to the research objectives','Thesis Table 4.8, checked against Tables 4.3–4.6 and §3.9',
        'Summarise what was demonstrated and where the evidence is limited. The architecture implements coordinated specialist analysis, but no ablation proves the minimum sufficient agent set. Accuracy, false positives and response time come from the original experiment. The 4.05-second mean is 0.05 seconds above the strict four-second warm target. Explanation readability is a descriptive manual rating, not a blinded user study. These are historical outcomes; the current system requires fresh performance evaluation.')
    cols=[.95,5.6,10.55]
    for x,w,label in zip(cols,[4.3,4.5,4.4],['RESEARCH AIM','REPORTED FINDING','INTERPRETATION']):
        s.text(x,2.93,w,.4,label,13,SAGE,MONO,1)
    rows=[
        ('Coordinate specialist analysis','Eight specialists + reporter','Implemented architecture'),
        ('Classify URLs accurately','94.7% accuracy','284 / 300 correct'),
        ('Reduce false alarms','2.0% false-positive rate','3 / 150 benign URLs flagged'),
        ('Provide timely feedback','4.05 s mean response','0.05 s above warm target'),
        ('Explain decisions clearly','9.0 / 10 readability','Descriptive manual rating'),
    ]
    for i,row in enumerate(rows):
        y=3.55+i*.78
        if i%2==0:s.rect(.8,y-.1,14.4,.69,PANEL,group=i+2)
        for j,(x,w,t) in enumerate(zip(cols,[4.3,4.5,4.4],row)):
            s.text(x,y,w,.52,t,18,ORANGE if j==1 else CREAM if j==0 else MUTED,group=i+2)
    s.text(.95,7.68,14,.4,'Historical evaluation findings; no new benchmark is claimed for the revised implementation.',16,MUTED,group=7)

    s=Slide('08 / Conclusion','Conclusion','CyberForge integrates threat analysis, supporting evidence and understandable warnings.','Thesis Chapter 5; revised Appendix C  /  Thank you — Questions',
        'Conclude with three contributions: an integrated specialist workflow; evidence-focused interpretation; and a current implementation with stronger scope, privacy and action boundaries. Historical results are promising on the selected set, but do not establish a benchmark for the revised source. Priorities: independent detector/error evaluation, controlled training/release and packaged/live verification. Invite questions on agent orchestration, evaluation limits, privacy, current MCP scope and training types. The original title remains unchanged; this is an updated system description, not fabricated experimental work.')
    s.text(.9,3.25,9.7,1.1,'94.7% accuracy  /  2.0% false alarms',30,ORANGE,HEAD,1)
    s.text(.95,4.25,9.5,.5,'Reported on the original 300-URL evaluation.',18,MUTED,group=1)
    for i,(t,b) in enumerate([('Contribution','Specialised analysis + structured evidence + usable explanations'),('Next validation','Independent data, per-detector errors and controlled comparisons'),('Next delivery','Verified releases, training isolation and bounded AI integrations')]):
        y=5.1+i*.85;s.text(.95,y,2.5,.4,t,17,SAGE,HEAD,2+i);s.text(3.55,y,7.4,.65,b,17,CREAM,group=2+i)
    s.picture(ASSETS/'orbit.gif',11.0,3.0,4.2,4.2,1)

def diagrams():
    # Word diagrams are deliberately separate from slide crops: readable labels,
    # high resolution, no presentation chrome or experimental claims.
    for filename,rows in [('architecture.png',[
        ('Clients','Tauri desktop | Rust CLI | grant-scoped MCP stdio | Expo companion'),
        ('Native boundary','Shared Rust services | OS keyring | consent | local ownership'),
        ('Backend','Node.js 24 / Express | verified account + workspace | device / agent workflows'),
        ('Analysis','Specialist DAG OR versioned detector endpoint | evidence + explicit coverage'),
        ('ML service','FastAPI on HF | boosted URL models | IsolationForest | BERT / Qwen | rules'),
        ('Data + intelligence','MongoDB | selected Appwrite integrations | local memory | configured OTX / caches')]),
        ('decision-flow.png',[
        ('1. Input','Manual URL or permitted observation; validate consent and redact sensitive fields'),
        ('2. Authority','Authenticate account; derive workspace and enforce operation scope'),
        ('3. Analyse','Select supported path; label unavailable engines and missing evidence'),
        ('4. Report','Return verdict + policy + provenance; advisory text is not execution'),
        ('5. Propose / approve','Only supported actions; explicit authority, target and expiry'),
        ('6. Execute / verify','Native permission or signed device command; retain actual execution receipt'),
        ('7. Retain context','Persist scoped evidence where allowed; memory is not weight training')])]:
        im=Image.new('RGB',(1800,50+230*len(rows)),'#'+CREAM);d=ImageDraw.Draw(im)
        for i,(title,body) in enumerate(rows):
            y=25+i*230;d.rounded_rectangle((40,y,1760,y+190),radius=18,fill='#'+BG)
            d.text((72,y+17),title,font=font(31,HEAD),fill='#'+ORANGE)
            for j,line in enumerate(wrapped(body,16.4,28)):
                d.text((72,y+78+j*48),line,font=font(28),fill='#'+CREAM)
            if i<len(rows)-1:d.line((900,y+194,900,y+220),fill='#'+SAGE,width=7);d.polygon([(887,y+210),(913,y+210),(900,y+228)],fill='#'+SAGE)
        im.save(ASSETS/filename)

def save():
    assert len(slides)==15
    for s in slides:s.animate()
    path=OUT/'CyberForge - Thesis Defence.pptx';prs.save(path)
    slides[0].im.save(OUT/'CyberForge - Slide Preview.pdf',save_all=True,append_images=[s.im for s in slides[1:]],resolution=100)
    thumbs=Image.new('RGB',(1600,5*328),'#'+INK);d=ImageDraw.Draw(thumbs)
    for i,s in enumerate(slides):
        thumb=s.im.resize((512,288),Image.Resampling.LANCZOS);x=(i%3)*533+10;y=(i//3)*328+8;thumbs.paste(thumb,(x,y));d.text((x,y+292),f'{i+1:02d}  '+all_notes[i]['title'].replace('\n',' ')[:62],font=font(10),fill='#'+CREAM)
    thumbs.save(OUT/'Slide overview.png')
    (OUT/'speaker-notes.json').write_text(json.dumps(all_notes,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'layout-audit.json').write_text(json.dumps(boxes,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Created',path,'with',len(slides),'slides and',sum(len(g) for s in slides for g in s.groups.values()),'animated shape builds.')

if __name__=='__main__':
    build();diagrams();save()
